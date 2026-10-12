from pathlib import Path
import json,re,pymupdf,html,hashlib
from docx import Document
from docx.shared import Inches,Pt

root=Path(__file__).resolve().parents[2]
pdf=pymupdf.open(r'C:\Users\lemin\Downloads\Reading v10.pdf')
questions=json.loads((root/'Backend/.work/inspection/questions-extracted.json').read_text(encoding='utf-8'))
answers5='B D C A A D B B C A D A D A D D B B A C D C C A B C B D C D'.split()
answers6='B C A A D A B C B C A D D A D A'.split()
answers7='C A A C B A C D C B D A C D D D C A D C B B B A C B B C A D A C B A B C D A B B A C B C C D C B A A C D B C'.split()
answers=answers5+answers6+answers7
assert len(questions)==len(answers)==100
for q,a in zip(questions,answers):
    q['answer']=a;q['status']='proposed'
    q['text']=q['text'].replace('------c','-------')
    if 131<=q['number']<=146:q['text']=str(q['number'])
    for k,v in q['options'].items():q['options'][k]=re.sub(r'\s+I$','',v.split(' Stop!')[0]).strip()
    if q['number']==151:q['options']['D']='A store manager'
groups=[(131,134,[(5,195,410)]),(135,138,[(6,85,514)]),(139,142,[(7,85,435)]),(143,146,[(8,85,324)]),(147,148,[(9,194,557)]),(149,150,[(10,80,335)]),(151,152,[(11,82,490)]),(153,154,[(12,85,410)]),(155,157,[(13,79,450)]),(158,160,[(14,81,400)]),(161,163,[(15,81,385)]),(164,167,[(16,80,490)]),(168,171,[(18,80,600)]),(172,175,[(20,80,368)]),(176,180,[(22,82,505),(22,510,718)]),(181,185,[(24,78,415),(24,429,713)]),(186,190,[(26,80,419),(26,430,659),(27,48,365)]),(191,195,[(28,80,450),(28,467,659),(29,42,253)]),(196,200,[(30,76,311),(30,322,664),(31,59,312)])]
part6={
131:'To: All Employees\nSubject: Server Maintenance Reminder\nDate: March 11\n\nPlease be advised it is time for the IT department to {{131}} mandatory server maintenance and updates. {{132}}. There will be no Internet service in the building from 7:00 P.M. Wednesday, March 12, until 9:00 A.M. Thursday, March 13. In addition, please note there will be no remote access available. Therefore, you will not be able to log in to the server from outside the office.\n\n{{133}} this regularly scheduled maintenance, you will not have access to your e-mail, calendar, or contacts. Employees should plan accordingly. We {{134}} any inconvenience this may cause.',
135:'22 April\n\nMs. Anna Schoorl\nRodezand 334\n3011 AV Rotterdam\nNetherlands\n\nDear Ms. Schoorl,\n\nCongratulations on your remarkable {{135}} in the Netherlands, Belgium, and Luxembourg. Your region has improved its on-time delivery performance for each of the past seven quarters. {{136}}.\n\nI am pleased to offer you a promotion to Director of European Operations. The position {{137}} in Hamburg, Germany. I realize that relocating may be difficult for you. {{138}}, I certainly hope that you will take time to consider this opportunity. Please call me at your earliest convenience so that we can discuss any concerns you may have.\n\nThank you for being a part of the Unocity Shipping family.\n\nSincerely,\nXia Hsu, Director of Operations\nUnocity Shipping, Inc.',
139:'To: Wu Investment Services employees\nFrom: Eileen Suen, Office Manager\nRe: Jacob Wu\nDate: 15 August\n\nTo All Staff,\n\nAs many of you are aware, Jacob Wu, our long-standing Chief Executive Officer, {{139}} on 1 October. Twenty years ago, Mr. Wu set out to create a Hong Kong-based services firm with an international scope. {{140}}. Wu Investment Services currently serves clients in seventeen countries, 95 percent of whom have chosen to invest with us for the long term.\n\nMr. Wu will be {{141}} by Thomas Wu, his son, who has served as Vice President of Wu Investment Services for the past four years.\n\nA gathering will be held on 28 September to celebrate Mr. Wu\'s {{142}} career. I will send further information about the event closer to the date. Thank you.\n\nEileen',
143:'From: Madeleine DeVries, Director of Operations\nTo: All Employees\nDate: June 1\nRe: Travel Policy\n\nTo help reduce {{143}}, the officers have voted to change the company\'s travel policy. The revised policy will be {{144}} on June 15. From that point forward, employees traveling within the country will be required to submit their travel requests to the accounting office no later than three weeks before the date of departure. {{145}}.\n\n{{146}} exceptions to this policy will be decided on a case-by-case basis and must first be approved by the individual employee\'s supervisor.'}
out=root/'Backend/docs';out.mkdir(exist_ok=True);assets=out/'reading-assets';assets.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(4)
def tag(s):
    p=doc.add_paragraph(s)
    p.paragraph_format.keep_with_next=s.startswith(('<question ', '<part ', '<group ', '<document '))
def escape(s):return html.escape(s,quote=False)
tag('<exam name="ETS TOEIC Reading Test 10" minutes="75" total="100" answerStatus="proposed">')
dataset={'name':'ETS TOEIC Reading Test 10','answerStatus':'proposed','sourcePdfSha256':hashlib.sha256(Path(r'C:\Users\lemin\Downloads\Reading v10.pdf').read_bytes()).hexdigest(),'questions':questions,'groups':[]}
def question(q):
    tag(f'<question number="{q["number"]}" answer="{q["answer"]}" status="proposed">')
    tag('<stem>'+escape(q['text'])+'</stem>')
    for k in 'ABCD':tag(f'<option letter="{k}">'+escape(q['options'][k])+'</option>')
    tag('</question>')
for part,lo,hi in [(5,101,130),(6,131,146),(7,147,200)]:
    if part>5:doc.add_page_break()
    tag(f'<part number="{part}">')
    if part==5:
        for q in questions[:30]:question(q)
    else:
        for first,last,regions in groups:
            if not lo<=first<=hi:continue
            tag(f'<group id="g{first}-{last}">');g={'key':f'g{first}-{last}','first':first,'last':last,'documents':[]}
            for order,(page,y1,y2) in enumerate(regions,1):
                p=pdf[page-1]
                if part==7:
                    # Reading documents end above question columns; exclude their headers/options.
                    stops=[]
                    for block in p.get_text('dict')['blocks']:
                        for line in block.get('lines',[]):
                            line_text=''.join(span['text'] for span in line['spans'])
                            match=re.match(r'\s*(\d{3})\.\s',line_text)
                            if match and first<=int(match[1])<=last and line['bbox'][1]>y1:
                                stops.append(line['bbox'][1]-12)
                    if stops:y2=min(y2,min(stops))
                rect=pymupdf.Rect(40,y1,515,y2);path=assets/f'g{first}-doc{order}.png';p.get_pixmap(matrix=pymupdf.Matrix(2,2),clip=rect).save(path)
                text=part6[first] if part==6 else p.get_text(clip=rect,sort=True).strip()
                tag(f'<document order="{order}">')
                xmltext=re.sub(r'\{\{(\d+)\}\}',r'<blank number="\1"/>',escape(text))
                tag('<text>'+xmltext+'</text>')
                para=doc.add_paragraph('<image ref="embedded"/>');para.add_run().add_break();para.add_run().add_picture(str(path),width=Inches(6.1))
                tag('</document>');g['documents'].append({'page':page,'crop':[40,y1,515,y2],'text':text,'image':str(path.relative_to(root)).replace('\\','/')})
            for q in questions:
                if first<=q['number']<=last:question(q)
            tag('</group>');dataset['groups'].append(g)
    tag('</part>')
tag('</exam>');doc.save(out/'TOEIC-Reading-Test10-Tagged.docx')
(out/'reading-test10.json').write_text(json.dumps(dataset,indent=2,ensure_ascii=False),encoding='utf-8')
(out/'ANSWER-REVIEW.md').write_text('# Bộ đáp án đề xuất\n\nKhông có bảng đáp án ETS chính thức trong PDF được cung cấp. Toàn bộ 100 câu đang ở trạng thái proposed. Đã giải dựa trên ngữ pháp và bài đọc; cần người dùng đối chiếu trước khi chốt bản nộp. Không đổi sang verified nếu chưa đối chiếu.\n\n| Câu | Đề xuất |\n|---|---|\n'+'\n'.join(f'| {q["number"]} | {q["answer"]} |' for q in questions),encoding='utf-8')
print('Generated 100 questions; 19 groups;',sum(len(g['documents']) for g in dataset['groups']),'documents; all answers proposed.')

